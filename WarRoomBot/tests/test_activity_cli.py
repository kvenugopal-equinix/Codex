"""A completed escalation retains its event before updating the optional workbook."""
import csv

from openpyxl import load_workbook

from warroom import cli
from warroom.activity_report import SHEET, EVENT_ID_COLUMN
from warroom.config import Settings
from warroom.dialer import DialDecision, DialResult, InviteOutcome


def _settings(tmp_path, destination):
    return Settings(dry_run=True, live_calling=False, target_mode="test", test_number="+14155550123",
                    allow_roster_calls=False, roster_path=tmp_path / "roster.xlsx",
                    join_log_path=tmp_path / "logs" / "join_log.csv", activity_workbook_path=destination)


def _fake_call(monkeypatch):
    monkeypatch.setattr(cli, "choose_target", lambda *_: DialDecision("", "TEST", False, "dry run"))
    monkeypatch.setattr(cli, "invite", lambda *_: InviteOutcome(""))
    monkeypatch.setattr(cli, "dial", lambda *_: DialResult("DRY-RUN: no call placed"))


def test_escalation_writes_local_event_then_selected_workbook(tmp_path, monkeypatch):
    _fake_call(monkeypatch)
    workbook = tmp_path / "chosen activity.xlsx"
    settings = _settings(tmp_path, workbook)
    cli._escalate(settings, "Visitor", [], None)
    with settings.join_log_path.open(newline="", encoding="utf-8") as stream:
        event = next(csv.DictReader(stream))
    sheet = load_workbook(workbook)[SHEET]
    assert sheet.cell(2, EVENT_ID_COLUMN).value == event["event_id"]
    assert sheet.cell(2, 5).value == "Visitor"


def test_unavailable_workbook_keeps_local_event_without_repeating_call(tmp_path, monkeypatch, capsys):
    _fake_call(monkeypatch)
    settings = _settings(tmp_path, tmp_path / "unavailable" / "chosen.xlsx")
    result, _ = cli._escalate(settings, "Visitor", [], None)
    assert result.text == "DRY-RUN: no call placed"
    with settings.join_log_path.open(newline="", encoding="utf-8") as stream:
        assert len(list(csv.DictReader(stream))) == 1
    assert "ACTIVITY REPORT PENDING" in capsys.readouterr().err
