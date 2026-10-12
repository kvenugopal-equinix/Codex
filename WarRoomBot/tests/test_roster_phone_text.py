"""Excel phone formatting preserves international prefixes without guessing destinations."""
import importlib.util
from pathlib import Path

from openpyxl import load_workbook
import pytest

from warroom.roster import _row_to_shift, load_roster


SCRIPT = Path(__file__).resolve().parents[1] / "tools" / "make_roster.py"
SPEC = importlib.util.spec_from_file_location("make_roster_phone_test", SCRIPT)
make_roster = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(make_roster)


def row(phone):
    return dict(name="Engineer", phone=phone, shift_start="00:00", shift_end="00:00", priority=1)


def test_new_template_phone_column_and_cells_are_text(tmp_path):
    path = tmp_path / "roster.xlsx"
    assert make_roster.main(["make_roster.py", str(path)]) == 0
    book = load_workbook(path)
    try:
        sheet = book["roster"]
        assert sheet.column_dimensions["C"].number_format == "@"
        assert all(sheet.cell(row=i, column=3).number_format == "@" for i in range(2, 5))
        sheet["A2"] = "Engineer"
        sheet["C2"] = "+6590000001"
        book.save(path)
    finally:
        book.close()
    assert load_roster(path)[0].phone == "+6590000001"


def test_unchanged_old_template_format_repaired_without_value_changes(tmp_path):
    path = tmp_path / "roster.xlsx"
    make_roster.main(["make_roster.py", str(path)])
    book = load_workbook(path)
    try:
        sheet = book["roster"]
        before = list(sheet.values)
        sheet.column_dimensions["C"].number_format = "General"
        for i in range(2, 5):
            sheet.cell(row=i, column=3).number_format = "General"
        book.save(path)
    finally:
        book.close()
    make_roster.main(["make_roster.py", str(path)])
    book = load_workbook(path)
    try:
        assert list(book["roster"].values) == before
        assert book["roster"].column_dimensions["C"].number_format == "@"
        assert book["roster"]["C2"].number_format == "@"
    finally:
        book.close()


def test_user_roster_is_not_rewritten(tmp_path):
    path = tmp_path / "roster.xlsx"
    make_roster.main(["make_roster.py", str(path)])
    book = load_workbook(path)
    try:
        book["roster"]["A2"] = "Actual engineer"
        book["roster"]["C2"] = 6590000001
        book["roster"]["C2"].number_format = "General"
        book.save(path)
    finally:
        book.close()
    before = path.read_bytes()
    make_roster.main(["make_roster.py", str(path)])
    assert path.read_bytes() == before
    with pytest.raises(ValueError, match="Format the Excel phone column as Text"):
        load_roster(path)


def test_unrecognized_existing_file_is_preserved(tmp_path):
    path = tmp_path / "roster.xlsx"
    data = b"not a generated workbook"
    path.write_bytes(data)
    assert make_roster.main(["make_roster.py", str(path)]) == 0
    assert path.read_bytes() == data


@pytest.mark.parametrize("phone", [6590000001, 6590000001.0, "6590000001", " 6590000001 ", True])
def test_numeric_phone_is_not_silently_assigned_a_country_code(phone):
    with pytest.raises(ValueError, match="re-enter.*country code"):
        _row_to_shift(row(phone))


def test_explicit_international_text_and_placeholder_still_load():
    assert _row_to_shift(row("  +6590000001  ")).phone == "+6590000001"
    assert _row_to_shift(row("+XXXXXXXXXXX")).is_placeholder


def test_csv_numeric_phone_is_rejected_with_same_guidance(tmp_path):
    path = tmp_path / "roster.csv"
    path.write_text("name,phone,shift_start,shift_end,priority\nEngineer,6590000001,00:00,00:00,1\n")
    with pytest.raises(ValueError, match="No country code is inferred"):
        load_roster(path)
