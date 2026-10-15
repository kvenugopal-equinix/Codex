"""Bot-created XLSX publication from the durable local join journal."""
import csv
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from warroom.activity_report import (
    COMMENTS_COLUMN, EVENT_ID_COLUMN, HEADERS, SHEET, WorkbookConflictError,
    publish_workbook,
)
from warroom.joinlog import FIELDS, log_join


def _records(path):
    with path.open(newline="", encoding="utf-8") as stream:
        return list(csv.DictReader(stream))


def _activity(path):
    return load_workbook(path)[SHEET]


def test_selection_creates_a_bot_workbook_before_first_event(tmp_path):
    journal = tmp_path / "local" / "joins.csv"
    workbook = tmp_path / "OneDrive O'Connor 東京" / "activity #1.xlsx"
    workbook.parent.mkdir()
    assert publish_workbook(journal, workbook) == 0
    assert workbook.is_file()
    assert not journal.exists()
    sheet = _activity(workbook)
    assert [cell.value for cell in sheet[1]] == HEADERS
    assert sheet.protection.sheet
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == f"A1:Q1"


def test_legacy_journal_is_given_stable_ids_before_backfill(tmp_path):
    journal, workbook = tmp_path / "joins.csv", tmp_path / "Activity.xlsx"
    with journal.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(FIELDS[:-1])
        writer.writerow(["" if field != "participant_name" else "Old Visitor" for field in FIELDS[:-1]])
    assert publish_workbook(journal, workbook) == 1
    event_id = _records(journal)[0]["event_id"]
    assert _activity(workbook).cell(2, EVENT_ID_COLUMN).value == event_id
    before = workbook.read_bytes()
    assert publish_workbook(journal, workbook) == 0
    assert workbook.read_bytes() == before


def test_backfill_and_retries_add_each_event_once_with_comments_by_event_id(tmp_path):
    journal = tmp_path / "joins.csv"
    workbook = tmp_path / "Activity.xlsx"
    first = log_join(journal, participant_name="Same Visitor", notes="bot note")
    second = log_join(journal, participant_name="Same Visitor")
    assert first["event_id"] != second["event_id"]
    assert publish_workbook(journal, workbook) == 2

    book = load_workbook(workbook)
    sheet = book[SHEET]
    assert sheet.cell(2, COMMENTS_COLUMN).protection.locked is False
    assert sheet.cell(2, EVENT_ID_COLUMN).protection.locked is True
    assert sheet.cell(2, FIELDS.index("notes") + 1).value == "bot note"
    sheet.cell(2, COMMENTS_COLUMN).value = "Comment on first join"
    # Move the entire data row, including its event ID and comment, as sorting would.
    for col in range(1, COMMENTS_COLUMN + 1):
        a, b = sheet.cell(2, col), sheet.cell(3, col)
        a.value, b.value = b.value, a.value
    book.save(workbook)

    third = log_join(journal, participant_name="Same Visitor")
    assert publish_workbook(journal, workbook) == 1
    assert publish_workbook(journal, workbook) == 0
    sheet = _activity(workbook)
    rows = {
        sheet.cell(row, EVENT_ID_COLUMN).value: sheet.cell(row, COMMENTS_COLUMN).value
        for row in range(2, sheet.max_row + 1)
    }
    assert rows == {first["event_id"]: "Comment on first join",
                    second["event_id"]: None, third["event_id"]: None}
    assert len(_records(journal)) == 3
    assert all(sheet.cell(row, COMMENTS_COLUMN).protection.locked is False
               for row in range(2, sheet.max_row + 1))


def test_formula_like_activity_stays_literal_and_plus_is_preserved(tmp_path):
    journal, workbook = tmp_path / "joins.csv", tmp_path / "Activity.xlsx"
    log_join(journal, participant_name="=HYPERLINK(\"https://example.test\")",
             notes="+14155550123")
    publish_workbook(journal, workbook)
    sheet = _activity(workbook)
    name = sheet.cell(2, FIELDS.index("participant_name") + 1)
    note = sheet.cell(2, FIELDS.index("notes") + 1)
    assert name.value.startswith("=HYPERLINK") and name.data_type == "s"
    assert note.value == "+14155550123" and note.data_type == "s"


def test_unrelated_and_malformed_destinations_are_left_untouched(tmp_path):
    journal, workbook = tmp_path / "joins.csv", tmp_path / "Activity.xlsx"
    log_join(journal, participant_name="Visitor")
    other = Workbook()
    other.active["A1"] = "Someone else's spreadsheet"
    other.save(workbook)
    before = workbook.read_bytes()
    with pytest.raises(WorkbookConflictError, match="not created by War Room Bot"):
        publish_workbook(journal, workbook)
    assert workbook.read_bytes() == before

    workbook.write_bytes(b"not an xlsx")
    before = workbook.read_bytes()
    with pytest.raises(WorkbookConflictError, match="not a readable"):
        publish_workbook(journal, workbook)
    assert workbook.read_bytes() == before


def test_changed_bot_cell_and_duplicate_id_refuse_publication(tmp_path):
    journal, workbook = tmp_path / "joins.csv", tmp_path / "Activity.xlsx"
    log_join(journal, participant_name="Visitor")
    publish_workbook(journal, workbook)
    book = load_workbook(workbook)
    book[SHEET].cell(2, FIELDS.index("participant_name") + 1).value = "Changed"
    book.save(workbook)
    before = workbook.read_bytes()
    with pytest.raises(WorkbookConflictError, match="differs"):
        publish_workbook(journal, workbook)
    assert workbook.read_bytes() == before

    book = load_workbook(workbook)
    sheet = book[SHEET]
    sheet.cell(2, FIELDS.index("participant_name") + 1).value = "Visitor"
    for col in range(1, COMMENTS_COLUMN + 1):
        sheet.cell(3, col).value = sheet.cell(2, col).value
    book.save(workbook)
    before = workbook.read_bytes()
    with pytest.raises(WorkbookConflictError, match="duplicate event ID"):
        publish_workbook(journal, workbook)
    assert workbook.read_bytes() == before


def test_local_edit_during_save_is_detected_and_journal_remains(tmp_path, monkeypatch):
    import warroom.activity_report as report

    journal, workbook = tmp_path / "joins.csv", tmp_path / "Activity.xlsx"
    first = log_join(journal, participant_name="First")
    publish_workbook(journal, workbook)
    second = log_join(journal, participant_name="Second")
    real_save = report._save_temporary

    def edit_while_saving(book, destination):
        temporary = real_save(book, destination)
        human_book = load_workbook(destination)
        human_book[SHEET].cell(2, COMMENTS_COLUMN).value = "Saved while bot rendered"
        human_book.save(destination)
        return temporary

    monkeypatch.setattr(report, "_save_temporary", edit_while_saving)
    with pytest.raises(WorkbookConflictError, match="changed during publication"):
        publish_workbook(journal, workbook)
    assert [row["event_id"] for row in _records(journal)] == [first["event_id"], second["event_id"]]
    sheet = _activity(workbook)
    assert sheet.cell(2, COMMENTS_COLUMN).value == "Saved while bot rendered"
    assert sheet.max_row == 2
    assert not list(tmp_path.glob(".Activity.xlsx.*.tmp"))


def test_missing_destination_folder_keeps_local_event_for_retry(tmp_path):
    journal = tmp_path / "joins.csv"
    event = log_join(journal, participant_name="Visitor")
    destination = tmp_path / "offline" / "Activity.xlsx"
    with pytest.raises(FileNotFoundError):
        publish_workbook(journal, destination)
    assert _records(journal)[0]["event_id"] == event["event_id"]
    destination.parent.mkdir()
    assert publish_workbook(journal, destination) == 1


def test_file_appearing_during_creation_is_not_overwritten(tmp_path, monkeypatch):
    import warroom.activity_report as report

    journal, workbook = tmp_path / "joins.csv", tmp_path / "Activity.xlsx"
    log_join(journal, participant_name="Visitor")

    def competing_creator(temporary, destination):
        Path(destination).write_bytes(b"created elsewhere")
        raise FileExistsError("already exists")

    monkeypatch.setattr(report.os, "link", competing_creator)
    with pytest.raises(WorkbookConflictError, match="appeared during creation"):
        publish_workbook(journal, workbook)
    assert workbook.read_bytes() == b"created elsewhere"
    assert len(_records(journal)) == 1
    assert not list(tmp_path.glob(".Activity.xlsx.*.tmp"))


def test_local_excel_lock_file_blocks_update_without_losing_event(tmp_path):
    journal, workbook = tmp_path / "joins.csv", tmp_path / "Activity.xlsx"
    log_join(journal, participant_name="First")
    publish_workbook(journal, workbook)
    second = log_join(journal, participant_name="Second")
    lock = tmp_path / "~$Activity.xlsx"
    lock.write_bytes(b"Excel")
    before = workbook.read_bytes()
    with pytest.raises(WorkbookConflictError, match="open in Excel"):
        publish_workbook(journal, workbook)
    assert workbook.read_bytes() == before
    lock.unlink()
    assert publish_workbook(journal, workbook) == 1
    assert _activity(workbook).cell(3, EVENT_ID_COLUMN).value == second["event_id"]
