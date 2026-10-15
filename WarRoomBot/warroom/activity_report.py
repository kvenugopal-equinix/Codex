"""Publish the local join journal to a bot-created, commentable Excel workbook.

The CSV remains the durable source of activity. This module can retry publication
after a destination becomes available without making another call or adding a
duplicate row. A local XLSX rewrite is not an Excel/OneDrive coauthoring operation.
"""

import argparse
import csv
import errno
import io
import os
import re
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter

from .joinlog import FIELDS, _locked, _migrate


SHEET = "Activity"
META_SHEET = "__WarRoomBot"
MARKER = "warroom-activity-workbook-v1"
HEADERS = FIELDS + ["Comments"]
COMMENTS_COLUMN = len(HEADERS)
EVENT_ID_COLUMN = FIELDS.index("event_id") + 1
_INVALID_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f]")


class WorkbookConflictError(ValueError):
    """The selected workbook cannot safely be reconciled with the local journal."""


def _events(local_log_path: Path) -> list:
    """Read a consistent journal snapshot while the caller holds its writer lock."""
    if not local_log_path.exists() or local_log_path.stat().st_size == 0:
        return []
    if _migrate(local_log_path) != FIELDS:
        raise ValueError("local join log has an unrecognized header; workbook was not changed")
    with local_log_path.open(newline="", encoding="utf-8") as stream:
        reader = csv.DictReader(stream)
        if reader.fieldnames != FIELDS:
            raise ValueError("local join log header changed; workbook was not changed")
        events = []
        ids = set()
        for number, row in enumerate(reader, start=2):
            if None in row or any(value is None for value in row.values()):
                raise ValueError(f"local join log row {number} is malformed; workbook was not changed")
            event_id = row["event_id"]
            if not event_id or event_id in ids:
                raise ValueError(f"local join log row {number} has a missing or duplicate event ID")
            ids.add(event_id)
            events.append(row)
        return events


def _new_workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(COMMENTS_COLUMN)}1"
    sheet.sheet_view.showGridLines = False
    header_fill = PatternFill(fill_type="solid", fgColor="17365D")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 32
    for index, field in enumerate(HEADERS, start=1):
        width = 42 if field in ("notes", "Comments") else 25
        if field in ("participant_name", "engineer_name"):
            width = 30
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.protection.sheet = True
    sheet.protection.autoFilter = False  # permit use of the existing filter
    sheet.protection.selectLockedCells = False
    sheet.protection.selectUnlockedCells = False

    meta = book.create_sheet(META_SHEET)
    meta["A1"] = MARKER
    meta.sheet_state = "veryHidden"
    return book, sheet


def _read_owned_workbook(original: bytes):
    try:
        book = load_workbook(io.BytesIO(original))
    except Exception as exc:
        raise WorkbookConflictError("selected file is not a readable War Room Bot workbook") from exc
    if len(book.sheetnames) != 2 or set(book.sheetnames) != {SHEET, META_SHEET} or book[META_SHEET]["A1"].value != MARKER:
        raise WorkbookConflictError("selected file was not created by War Room Bot")
    sheet = book[SHEET]
    if [sheet.cell(1, index).value for index in range(1, COMMENTS_COLUMN + 1)] != HEADERS:
        raise WorkbookConflictError("activity workbook columns changed; file was not changed")
    if sheet.max_column > COMMENTS_COLUMN or not sheet.protection.sheet:
        raise WorkbookConflictError("activity workbook structure changed; file was not changed")
    if any(not cell.protection.locked for cell in sheet[1]):
        raise WorkbookConflictError("activity workbook header protection changed; file was not changed")
    return book, sheet


def _cell_text(value: str) -> str:
    """Keep source text literal, including leading +, =, @ and hyphens."""
    if len(value) > 32767:
        raise ValueError("an activity value exceeds Excel's 32,767-character cell limit")
    if _INVALID_XML.search(value):
        raise ValueError("an activity value contains a character Excel cannot store")
    return value


def _append(sheet, event: dict) -> None:
    number = sheet.max_row + 1
    for index, field in enumerate(FIELDS, start=1):
        cell = sheet.cell(number, index)
        cell.value = _cell_text(event[field])
        cell.data_type = "s"  # openpyxl otherwise interprets leading '=' as a formula
        cell.protection = Protection(locked=True)
        cell.alignment = Alignment(vertical="top", wrap_text=field == "notes")
    comment = sheet.cell(number, COMMENTS_COLUMN)
    comment.protection = Protection(locked=False)
    comment.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.auto_filter.ref = f"A1:{get_column_letter(COMMENTS_COLUMN)}{number}"


def _check_existing_rows(sheet, journal: dict) -> set:
    seen = set()
    for number in range(2, sheet.max_row + 1):
        event_id = sheet.cell(number, EVENT_ID_COLUMN).value
        if event_id not in journal or event_id in seen:
            raise WorkbookConflictError(
                f"activity workbook row {number} has an unknown or duplicate event ID; file was not changed"
            )
        seen.add(event_id)
        event = journal[event_id]
        for index, field in enumerate(FIELDS, start=1):
            cell = sheet.cell(number, index)
            if not cell.protection.locked:
                raise WorkbookConflictError(
                    f"activity workbook row {number} has an unlocked bot cell; file was not changed"
                )
            actual = cell.value
            # Excel stores a bot-authored empty string as an empty cell.
            if actual is None:
                actual = ""
            if actual != _cell_text(event[field]):
                raise WorkbookConflictError(
                    f"activity workbook row {number} differs from its local event; file was not changed"
                )
        if sheet.cell(number, COMMENTS_COLUMN).protection.locked:
            raise WorkbookConflictError(
                f"activity workbook row {number} has a locked Comments cell; file was not changed"
            )
    return seen


def _save_temporary(book, destination: Path) -> Path:
    fd, temporary = tempfile.mkstemp(prefix=f".{destination.name}.", suffix=".tmp", dir=destination.parent)
    os.close(fd)
    path = Path(temporary)
    try:
        book.save(path)
        with path.open("r+b") as stream:
            os.fsync(stream.fileno())
    except BaseException:
        path.unlink(missing_ok=True)
        raise
    return path


def _rename_exclusive_macos(temporary: Path, destination: Path) -> None:
    """Use macOS's no-replace rename when File Provider disallows hard links."""
    import ctypes

    libc = ctypes.CDLL(None, use_errno=True)
    try:
        renamex = libc.renamex_np
    except AttributeError as exc:
        raise OSError(errno.ENOSYS, "exclusive rename is unavailable", str(destination)) from exc
    renamex.argtypes = (ctypes.c_char_p, ctypes.c_char_p, ctypes.c_uint)
    renamex.restype = ctypes.c_int
    if renamex(os.fsencode(temporary), os.fsencode(destination), 0x00000004) != 0:  # RENAME_EXCL
        error = ctypes.get_errno()
        if error == errno.EEXIST:
            raise WorkbookConflictError("selected workbook appeared during creation; file was not changed")
        raise OSError(error, os.strerror(error), str(destination))


def _publish_new(temporary: Path, destination: Path) -> None:
    """Link into place only if no file exists; never replace an unrelated file."""
    try:
        os.link(temporary, destination)
    except FileExistsError as exc:
        raise WorkbookConflictError("selected workbook appeared during creation; file was not changed") from exc
    except OSError:
        if sys.platform != "darwin":
            raise
        _rename_exclusive_macos(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)


def _publish_update(temporary: Path, destination: Path, original: bytes) -> None:
    try:
        if (destination.parent / ("~$" + destination.name)).exists():
            raise WorkbookConflictError("selected workbook appears open in Excel; retry after closing it")
        if not destination.exists() or destination.read_bytes() != original:
            raise WorkbookConflictError("selected workbook changed during publication; retry after it is saved")
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)


def publish_workbook(local_log_path: Path, workbook_path: Path) -> int:
    """Create/update the selected workbook from durable events; return rows added.

    Existing Comments are preserved by event ID. This refuses unrelated workbooks,
    changed bot values, Excel's local lock file, and a local edit detected while
    saving. It cannot detect an edit still unsynced on another device.
    """
    local_log_path = Path(local_log_path)
    destination = Path(workbook_path)
    if not destination.is_absolute() or destination.suffix.lower() != ".xlsx":
        raise ValueError("activity workbook must be an absolute .xlsx file path")
    if not destination.parent.is_dir():
        raise FileNotFoundError(f"activity workbook folder does not exist: {destination.parent}")
    if destination.is_symlink():
        raise WorkbookConflictError("selected workbook is a symbolic link; file was not changed")
    if (destination.parent / ("~$" + destination.name)).exists():
        raise WorkbookConflictError("selected workbook appears open in Excel; retry after closing it")

    local_log_path.parent.mkdir(parents=True, exist_ok=True)
    with _locked(local_log_path):
        events = _events(local_log_path)
        journal = {event["event_id"]: event for event in events}
        if destination.exists():
            original = destination.read_bytes()
            book, sheet = _read_owned_workbook(original)
            seen = _check_existing_rows(sheet, journal)
        else:
            original = None
            book, sheet = _new_workbook()
            seen = set()

        missing = [event for event in events if event["event_id"] not in seen]
        if original is not None and not missing:
            return 0
        for event in missing:
            _append(sheet, event)
        temporary = _save_temporary(book, destination)
        if original is None:
            _publish_new(temporary, destination)
        else:
            _publish_update(temporary, destination, original)
        return len(missing)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Create or update a War Room Bot activity workbook")
    parser.add_argument("local_log_path", type=Path)
    parser.add_argument("workbook_path", type=Path)
    args = parser.parse_args(argv)
    try:
        added = publish_workbook(args.local_log_path, args.workbook_path)
    except (OSError, ValueError) as exc:
        parser.exit(1, f"Activity workbook was not updated: {exc}\n")
    print(f"Activity workbook ready: {args.workbook_path} ({added} event(s) added)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
