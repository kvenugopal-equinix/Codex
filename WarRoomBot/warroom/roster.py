"""Roster loading and on-duty lookup (time-zone aware, overnight shifts, priority)."""
import csv
from dataclasses import dataclass
from datetime import datetime, time
from pathlib import Path
from zoneinfo import ZoneInfo

DAYS = ["mon", "tue", "wed", "thu", "fri", "sat", "sun"]


@dataclass(frozen=True)
class Shift:
    name: str
    zoom_email: str
    phone: str
    timezone: str
    days: frozenset
    start: time
    end: time
    priority: int
    active: bool
    notes: str = ""

    @property
    def is_placeholder(self) -> bool:
        return self.name.startswith("<") or "X" in self.phone.replace("+", "")


def parse_days(text: str) -> frozenset:
    t = (text or "").strip().lower()
    if t in ("", "daily", "all", "*"):
        return frozenset(range(7))
    out = set()
    for part in t.replace(" ", "").split(","):
        if "-" in part:
            a, b = (DAYS.index(p[:3]) for p in part.split("-"))
            i = a
            while True:
                out.add(i)
                if i == b:
                    break
                i = (i + 1) % 7
        else:
            out.add(DAYS.index(part[:3]))
    return frozenset(out)


def parse_time(v) -> time:
    if isinstance(v, time):
        return v
    if isinstance(v, datetime):
        return v.time()
    h, m = str(v).strip().split(":")[:2]
    return time(int(h), int(m))


def _phone_text(value) -> str:
    text = str(value if value is not None else "").strip()
    if isinstance(value, (int, float)) or (text and text.isdecimal()):
        raise ValueError("roster phone is numeric or missing its '+' prefix. Format the Excel phone column as Text "
                         "and re-enter +<country code><number>; changing the cell format alone does not restore '+'. "
                         "No country code is inferred")
    return text


def _row_to_shift(r: dict) -> Shift:
    g = lambda k: r.get(k) if r.get(k) is not None else ""
    return Shift(
        name=str(g("name")).strip(),
        zoom_email=str(g("zoom_email")).strip(),
        phone=_phone_text(g("phone")),
        timezone=str(g("timezone")).strip() or "UTC",
        days=parse_days(str(g("days"))),
        start=parse_time(g("shift_start")),
        end=parse_time(g("shift_end")),
        priority=int(g("priority") if g("priority") != "" else 99),
        active=str(g("active")).strip().lower() in ("yes", "y", "true", "1", ""),
        notes=str(g("notes")).strip(),
    )


def load_roster(path: Path) -> list:
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with open(path, newline="") as f:
            rows = list(csv.DictReader(f))
    else:
        from openpyxl import load_workbook
        book = load_workbook(path, data_only=True, read_only=True)
        try:
            ws = book["roster"]
            it = ws.iter_rows(values_only=True)
            hdr = [str(h).strip() for h in next(it)]
            rows = [dict(zip(hdr, r)) for r in it if r and any(c is not None for c in r)]
        finally:
            book.close()
    return [_row_to_shift(r) for r in rows]


def on_duty(shifts: list, now_utc: datetime, include_placeholders: bool = False) -> list:
    """Active shifts covering `now_utc`, best priority first (ties keep roster order)."""
    hits = []
    for s in shifts:
        if not s.active or (s.is_placeholder and not include_placeholders):
            continue
        local = now_utc.astimezone(ZoneInfo(s.timezone))
        t, wd = local.time(), local.weekday()
        if s.start < s.end:
            ok = s.start <= t < s.end and wd in s.days
        elif s.start > s.end:  # overnight: shift belongs to the day it starts on
            ok = (t >= s.start and wd in s.days) or (t < s.end and (wd - 1) % 7 in s.days)
        else:  # start == end: treat as 24h
            ok = wd in s.days
        if ok:
            hits.append(s)
    return sorted(hits, key=lambda s: s.priority)
